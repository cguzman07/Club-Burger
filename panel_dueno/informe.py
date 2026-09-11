"""Informe auxiliar de ingresos para el contador (Excel e impresión).

No es factura electrónica DIAN ni libro oficial. Es el soporte interno del POS.
"""

from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

BOGOTA = timezone(timedelta(hours=-5), name="America/Bogota")
ROOT = Path(__file__).resolve().parent
EMITIDOS = ROOT / "informes_emitidos.json"

EMPRESA_DEFAULT = {
    "razon_social": "Club Burger",
    "nombre_comercial": "Club Burger",
    "nit": "",
    "dv": "",
    "direccion": "",
    "ciudad": "",
    "departamento": "",
    "telefono": "",
    "correo": "",
    "responsabilidad_tributaria": "Pendiente de registrar ante la DIAN",
    "actividad_economica": "5611 — Expendio a la mesa de comidas preparadas",
    "representante_legal": "",
    "contador": "",
    "tarjeta_profesional": "",
}

TITULOS = {
    "dia": "Informe diario de ingresos por ventas",
    "semana": "Informe semanal de ingresos por ventas",
    "mes": "Informe mensual de ingresos por ventas",
    "anio": "Informe anual de ingresos por ventas — período gravable",
    "rango": "Informe de ingresos por ventas — rango de fechas",
}

MARCO_LEGAL = [
    "Naturaleza del documento. Este informe es un soporte auxiliar interno generado por el sistema de punto de venta (POS) Club Burger. Recoge los ingresos registrados en el período indicado, en pesos colombianos (COP).",
    "No sustituye la factura electrónica. No es factura de venta, nota débito, nota crédito ni documento equivalente ante la DIAN. Tampoco reemplaza la facturación electrónica ni los documentos previstos en el Estatuto Tributario y la Resolución DIAN vigente.",
    "No sustituye los libros oficiales. No reemplaza los libros de contabilidad ni los registros que el comerciante debe llevar conforme al Código de Comercio (arts. 48 y siguientes) y a las normas de información financiera aplicables.",
    "Uso para el contador. Se entrega como papeles de trabajo y base de conciliación entre el POS, la contabilidad y la facturación electrónica, para declaración de renta, IVA u otras obligaciones, según la responsabilidad tributaria del contribuyente.",
    "Impuestos. Los valores son ingresos brutos registrados en caja. El IVA, INC u otros impuestos, retenciones y bases gravables se determinan con la facturación electrónica y la calificación ante la DIAN, no con este archivo.",
    "Integridad. El consecutivo y la huella digital identifican esta emisión. Cualquier cambio posterior en el POS no altera este ejemplar. Conserve el archivo junto con los soportes de facturación electrónica.",
]


def empresa_desde_config(config: dict[str, Any] | None) -> dict[str, str]:
    datos = dict(EMPRESA_DEFAULT)
    crudo = (config or {}).get("empresa") or {}
    for clave, valor in crudo.items():
        if clave in datos and valor is not None:
            datos[clave] = str(valor).strip()
    return datos


def nit_texto(empresa: dict[str, str]) -> str:
    nit = (empresa.get("nit") or "").strip()
    dv = (empresa.get("dv") or "").strip()
    if not nit:
        return "Por registrar"
    return f"{nit}-{dv}" if dv else nit


def titulo_informe(periodo: str) -> str:
    clave = "anio" if periodo in {"año", "ano", "year"} else (periodo or "dia")
    return TITULOS.get(clave, TITULOS["rango"])


def fecha_latino(iso: str | None) -> str:
    if not iso:
        return "—"
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(iso))
    if not m:
        return str(iso)
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"


def _chunk(n: int) -> str:
    unidades = ("", "un", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve")
    diez = (
        "diez",
        "once",
        "doce",
        "trece",
        "catorce",
        "quince",
        "dieciséis",
        "diecisiete",
        "dieciocho",
        "diecinueve",
    )
    decenas = ("", "", "veinte", "treinta", "cuarenta", "cincuenta", "sesenta", "setenta", "ochenta", "noventa")
    centenas = ("", "ciento", "doscientos", "trescientos", "cuatrocientos", "quinientos", "seiscientos", "setecientos", "ochocientos", "novecientos")
    if n == 0:
        return ""
    if n == 100:
        return "cien"
    c, resto = divmod(n, 100)
    texto = centenas[c]
    if resto == 0:
        return texto
    if 10 <= resto <= 19:
        union = f"{texto} " if texto else ""
        return f"{union}{diez[resto - 10]}".strip()
    d, u = divmod(resto, 10)
    partes = [texto] if texto else []
    if d == 2 and u:
        partes.append(f"veinti{unidades[u]}" if u != 1 else "veintiún")
    elif d:
        if u:
            partes.append(f"{decenas[d]} y {unidades[u]}")
        else:
            partes.append(decenas[d])
    elif u:
        partes.append(unidades[u])
    return " ".join(p for p in partes if p).strip()


def numero_en_letras(valor: float | int) -> str:
    n = int(round(float(valor or 0)))
    if n < 0:
        return "MENOS " + numero_en_letras(-n)
    if n == 0:
        return "cero"
    if n == 1:
        return "un"
    millones, resto = divmod(n, 1_000_000)
    miles, unidades = divmod(resto, 1000)
    partes = []
    if millones == 1:
        partes.append("un millón")
    elif millones:
        partes.append(f"{_chunk(millones)} millones")
    if miles == 1:
        partes.append("mil")
    elif miles:
        pieza = _chunk(miles)
        if pieza.endswith(" veintiún"):
            pieza = pieza[:-8] + "veintiún"
        if pieza.endswith(" un"):
            pieza = pieza[:-3] + " un"
        partes.append(f"{pieza} mil")
    if unidades:
        partes.append(_chunk(unidades))
    texto = " ".join(partes).replace("  ", " ").strip()
    return texto.replace("un mil", "mil")


def pesos_mcte(valor: float | int) -> str:
    n = int(round(float(valor or 0)))
    letras = numero_en_letras(n).upper()
    if n == 1:
        return "UN PESO M/CTE"
    return f"{letras} PESOS M/CTE"


def siguiente_consecutivo() -> str:
    anio = datetime.now(BOGOTA).year
    data = {"anio": anio, "seq": 0}
    if EMITIDOS.exists():
        try:
            data = json.loads(EMITIDOS.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            pass
    if int(data.get("anio") or 0) != anio:
        data = {"anio": anio, "seq": 0}
    data["seq"] = int(data.get("seq") or 0) + 1
    EMITIDOS.write_text(json.dumps(data, indent=2), encoding="utf-8")
    return f"CB-IV-{anio}-{data['seq']:04d}"


def huella(payload: dict[str, Any]) -> str:
    canon = json.dumps(payload, ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(canon.encode("utf-8")).hexdigest()


def metadatos_informe(reporte: dict[str, Any], empresa: dict[str, str], usuario: str = "") -> dict[str, Any]:
    ahora = datetime.now(BOGOTA)
    periodo = str(reporte.get("periodo") or "dia")
    cuerpo = {
        "desde": reporte.get("desde"),
        "hasta": reporte.get("hasta"),
        "total": round(float(reporte.get("total") or 0), 2),
        "items": int(reporte.get("items") or 0),
        "lineas": int(reporte.get("lineas") or 0),
        "por_metodo": reporte.get("por_metodo") or {},
        "nit": nit_texto(empresa),
    }
    return {
        "consecutivo": siguiente_consecutivo(),
        "titulo": titulo_informe(periodo),
        "periodo": periodo,
        "etiqueta": reporte.get("etiqueta") or "",
        "desde": reporte.get("desde"),
        "hasta": reporte.get("hasta"),
        "desde_latino": fecha_latino(str(reporte.get("desde") or "")),
        "hasta_latino": fecha_latino(str(reporte.get("hasta") or "")),
        "generado": ahora.strftime("%d/%m/%Y %H:%M:%S"),
        "ciudad_fecha": f"{empresa.get('ciudad') or 'Colombia'}, {ahora.strftime('%d/%m/%Y')}",
        "fuente": "Libro de ventas en la nube" if reporte.get("fuente") == "nube" else "Registro local del POS",
        "moneda": "Pesos colombianos (COP)",
        "total": round(float(reporte.get("total") or 0), 2),
        "total_letras": pesos_mcte(reporte.get("total") or 0),
        "items": int(reporte.get("items") or 0),
        "lineas": int(reporte.get("lineas") or 0),
        "efectivo": round(float(reporte.get("efectivo") or 0), 2),
        "otros_medios": round(float(reporte.get("otros_medios") or 0), 2),
        "por_metodo": reporte.get("por_metodo") or {},
        "por_dia": reporte.get("por_dia") or [],
        "por_producto": reporte.get("por_producto") or [],
        "aviso": reporte.get("aviso"),
        "huella": huella(cuerpo),
        "usuario": usuario or "Administrador",
        "empresa": empresa,
        "nit": nit_texto(empresa),
        "nit_pendiente": not bool((empresa.get("nit") or "").strip()),
        "marco_legal": MARCO_LEGAL,
    }


def nombre_archivo(meta: dict[str, Any], ext: str) -> str:
    desde = str(meta.get("desde") or "").replace("-", "")
    hasta = str(meta.get("hasta") or "").replace("-", "")
    return f"ClubBurger_InformeIngresos_{meta.get('consecutivo')}_{desde}_{hasta}.{ext}"


def _borde() -> Border:
    s = Side(style="thin", color="1F2A44")
    return Border(left=s, right=s, top=s, bottom=s)


def _llenar(ws, fila: int, col: int, valor, font=None, fill=None, align=None, formato=None, merge=None):
    cell = ws.cell(fila, col, valor)
    cell.border = _borde()
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = align
    if formato:
        cell.number_format = formato
    if merge:
        ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=merge)
    return cell


def _pagina(ws, landscape: bool = False) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.6, right=0.6, top=0.7, bottom=0.7, header=0.3, footer=0.3)
    ws.oddHeader.left.text = "Club Burger · Informe auxiliar de ingresos"
    ws.oddFooter.left.text = "Soporte interno POS · No es factura DIAN"
    ws.oddFooter.right.text = "Página &P de &N"
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def construir_excel(meta: dict[str, Any]) -> bytes:
    empresa = meta["empresa"]
    navy = Font(name="Calibri", size=14, bold=True, color="1F2A44")
    titulo = Font(name="Calibri", size=16, bold=True, color="1F2A44")
    seccion = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    etiqueta = Font(name="Calibri", size=10, bold=True, color="1F2A44")
    cuerpo = Font(name="Calibri", size=10, color="1F2A44")
    mini = Font(name="Calibri", size=8, italic=True, color="5B6375")
    fill_navy = PatternFill("solid", fgColor="1F2A44")
    fill_gold = PatternFill("solid", fgColor="C4A35A")
    fill_gray = PatternFill("solid", fgColor="F4F1EA")
    fill_warn = PatternFill("solid", fgColor="F8E6C8")
    izq = Alignment(horizontal="left", vertical="center", wrap_text=True)
    der = Alignment(horizontal="right", vertical="center")
    cen = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cop = '"$"#,##0'

    wb = Workbook()

    car = wb.active
    car.title = "Carátula"
    car.sheet_properties.tabColor = "1F2A44"
    for col, ancho in enumerate((18, 22, 22, 22, 22, 18), 1):
        car.column_dimensions[get_column_letter(col)].width = ancho
    _llenar(car, 1, 1, empresa.get("nombre_comercial") or "Club Burger", navy, fill_gold, izq, merge=6)
    _llenar(car, 2, 1, meta["titulo"].upper(), titulo, fill_gray, izq, merge=6)
    _llenar(car, 3, 1, "Soporte auxiliar del sistema POS para papeles de trabajo del contador", mini, fill_gray, izq, merge=6)
    filas_id = [
        ("Consecutivo", meta["consecutivo"]),
        ("Razón social", empresa.get("razon_social") or "—"),
        ("Nombre comercial", empresa.get("nombre_comercial") or "—"),
        ("NIT", meta["nit"]),
        ("Dirección", empresa.get("direccion") or "—"),
        ("Ciudad / departamento", " / ".join(p for p in (empresa.get("ciudad"), empresa.get("departamento")) if p) or "—"),
        ("Teléfono", empresa.get("telefono") or "—"),
        ("Correo", empresa.get("correo") or "—"),
        ("Responsabilidad tributaria", empresa.get("responsabilidad_tributaria") or "—"),
        ("Actividad económica (CIIU)", empresa.get("actividad_economica") or "—"),
        ("Período del informe", meta["etiqueta"]),
        ("Desde", meta["desde_latino"]),
        ("Hasta", meta["hasta_latino"]),
        ("Moneda", meta["moneda"]),
        ("Fuente de datos", meta["fuente"]),
        ("Fecha y hora de generación", f"{meta['generado']} (America/Bogota)"),
        ("Elaborado por", meta["usuario"]),
        ("Huella SHA-256", meta["huella"]),
    ]
    for i, (k, v) in enumerate(filas_id, 5):
        _llenar(car, i, 1, k, etiqueta, fill_gray, izq, merge=2)
        _llenar(car, i, 3, v, cuerpo, None, izq, merge=6)
    r = 5 + len(filas_id) + 1
    _llenar(car, r, 1, "TOTAL INGRESOS DEL PERÍODO", seccion, fill_navy, izq, merge=3)
    _llenar(car, r, 4, meta["total"], Font(name="Calibri", size=14, bold=True, color="FFFFFF"), fill_navy, der, cop, merge=6)
    r += 1
    _llenar(car, r, 1, meta["total_letras"], cuerpo, fill_gold, izq, merge=6)
    r += 2
    if meta["nit_pendiente"]:
        _llenar(
            car,
            r,
            1,
            "AVISO: registre el NIT en panel_dueno/config.json (empresa.nit) antes de entregar este informe al contador.",
            Font(name="Calibri", size=9, bold=True, color="7A4A00"),
            fill_warn,
            izq,
            merge=6,
        )
        r += 2
    _llenar(car, r, 1, "Declaración de uso", seccion, fill_navy, izq, merge=6)
    r += 1
    car.row_dimensions[r].height = 72
    _llenar(
        car,
        r,
        1,
        "Este documento no es factura electrónica ni libro oficial de comercio. Es un informe auxiliar de los ingresos capturados por el POS, destinado al contador del negocio para conciliar y preparar obligaciones tributarias.",
        mini,
        None,
        izq,
        merge=6,
    )
    r += 3
    for col, texto in enumerate(("Elaboró", "Representante legal", "Contador público"), 1):
        c = 1 + (col - 1) * 2
        _llenar(car, r, c, texto, etiqueta, fill_gray, cen, merge=c + 1)
        _llenar(car, r + 1, c, "", cuerpo, None, cen, merge=c + 1)
        car.row_dimensions[r + 1].height = 36
        firma = [meta["usuario"], empresa.get("representante_legal") or "Nombre y firma", empresa.get("contador") or "Nombre, firma y T.P."][col - 1]
        _llenar(car, r + 2, c, firma, mini, None, cen, merge=c + 1)
    _pagina(car)

    res = wb.create_sheet("Resumen")
    for col, ancho in enumerate((36, 18, 18, 22), 1):
        res.column_dimensions[get_column_letter(col)].width = ancho
    _llenar(res, 1, 1, "Resumen del período", navy, fill_gold, izq, merge=4)
    _llenar(res, 2, 1, f"{meta['desde_latino']} a {meta['hasta_latino']}", mini, None, izq, merge=4)
    headers = ["Concepto", "Cantidad", "Valor (COP)", "Participación"]
    for i, h in enumerate(headers, 1):
        _llenar(res, 4, i, h, seccion, fill_navy, cen)
    total = float(meta["total"] or 0) or 1
    filas_r = [
        ("Ingresos brutos registrados", meta["items"], meta["total"]),
        ("Ventas en efectivo", None, meta["efectivo"]),
        ("Ventas por medios electrónicos y otros", None, meta["otros_medios"]),
        ("Líneas de venta (ítems detallados)", meta["lineas"], None),
    ]
    for i, (concepto, cant, valor) in enumerate(filas_r, 5):
        _llenar(res, i, 1, concepto, cuerpo, None, izq)
        _llenar(res, i, 2, cant if cant is not None else "—", cuerpo, None, der)
        if valor is None:
            _llenar(res, i, 3, "—", cuerpo, None, der)
            _llenar(res, i, 4, "—", cuerpo, None, der)
        else:
            _llenar(res, i, 3, float(valor), cuerpo, None, der, cop)
            _llenar(res, i, 4, float(valor) / total, cuerpo, None, der, "0.0%")
    _llenar(res, 9, 1, "Valor en letras", etiqueta, fill_gray, izq)
    _llenar(res, 9, 2, meta["total_letras"], cuerpo, None, izq, merge=4)
    _pagina(res)

    med = wb.create_sheet("Medios de pago")
    for col, ancho in enumerate((28, 18, 16, 18), 1):
        med.column_dimensions[get_column_letter(col)].width = ancho
    _llenar(med, 1, 1, "Ingresos por medio de pago", navy, fill_gold, izq, merge=4)
    for i, h in enumerate(("Medio de pago", "Valor (COP)", "% del total", "Observación"), 1):
        _llenar(med, 3, i, h, seccion, fill_navy, cen)
    metodos = meta["por_metodo"] or {}
    fila = 4
    if not metodos:
        _llenar(med, 4, 1, "Sin movimientos en el período", mini, None, izq, merge=4)
        fila = 5
    else:
        for nombre, valor in metodos.items():
            obs = "Caja menor / efectivo" if str(nombre).lower() == "efectivo" else "Medio electrónico o transferencia"
            _llenar(med, fila, 1, nombre, cuerpo, None, izq)
            _llenar(med, fila, 2, float(valor), cuerpo, None, der, cop)
            _llenar(med, fila, 3, float(valor) / total, cuerpo, None, der, "0.0%")
            _llenar(med, fila, 4, obs, mini, None, izq)
            fila += 1
    _llenar(med, fila, 1, "TOTAL", etiqueta, fill_gray, izq)
    _llenar(med, fila, 2, meta["total"], etiqueta, fill_gray, der, cop)
    _llenar(med, fila, 3, 1 if metodos else 0, etiqueta, fill_gray, der, "0.0%")
    _llenar(med, fila, 4, "", etiqueta, fill_gray, izq)
    _pagina(med)

    dia = wb.create_sheet("Libro diario")
    for col, ancho in enumerate((16, 16, 14, 18, 18), 1):
        dia.column_dimensions[get_column_letter(col)].width = ancho
    _llenar(dia, 1, 1, "Movimiento diario de ingresos", navy, fill_gold, izq, merge=5)
    _llenar(dia, 2, 1, "Libro auxiliar de ventas del POS · un renglón por día calendario", mini, None, izq, merge=5)
    for i, h in enumerate(("Fecha", "Fecha (ISO)", "Ítems", "Líneas", "Ingresos (COP)"), 1):
        _llenar(dia, 4, i, h, seccion, fill_navy, cen)
    fila = 5
    for item in meta["por_dia"]:
        iso = str(item.get("fecha") or "")
        _llenar(dia, fila, 1, fecha_latino(iso), cuerpo, None, cen)
        _llenar(dia, fila, 2, iso, mini, None, cen)
        _llenar(dia, fila, 3, int(item.get("items") or 0), cuerpo, None, der)
        _llenar(dia, fila, 4, int(item.get("lineas") or 0), cuerpo, None, der)
        _llenar(dia, fila, 5, float(item.get("total") or 0), cuerpo, None, der, cop)
        fila += 1
    if fila == 5:
        _llenar(dia, 5, 1, "Sin ventas en el período", mini, None, izq, merge=5)
        fila = 6
    _llenar(dia, fila, 1, "TOTAL PERÍODO", etiqueta, fill_gray, izq, merge=2)
    _llenar(dia, fila, 3, meta["items"], etiqueta, fill_gray, der)
    _llenar(dia, fila, 4, meta["lineas"], etiqueta, fill_gray, der)
    _llenar(dia, fila, 5, meta["total"], etiqueta, fill_gray, der, cop)
    dia.auto_filter.ref = f"A4:E{max(fila, 5)}"
    dia.freeze_panes = "A5"
    dia.oddHeader.left.text = f"{empresa.get('razon_social')} · NIT {meta['nit']}"
    _pagina(dia, landscape=True)

    prod = wb.create_sheet("Productos")
    for col, ancho in enumerate((8, 36, 14, 18, 16), 1):
        prod.column_dimensions[get_column_letter(col)].width = ancho
    _llenar(prod, 1, 1, "Ingresos por producto", navy, fill_gold, izq, merge=5)
    for i, h in enumerate(("#", "Producto", "Cantidad", "Ingresos (COP)", "%"), 1):
        _llenar(prod, 3, i, h, seccion, fill_navy, cen)
    fila = 4
    for i, item in enumerate(meta["por_producto"], 1):
        _llenar(prod, fila, 1, i, cuerpo, None, cen)
        _llenar(prod, fila, 2, item.get("nombre") or "Producto", cuerpo, None, izq)
        _llenar(prod, fila, 3, int(item.get("cantidad") or 0), cuerpo, None, der)
        _llenar(prod, fila, 4, float(item.get("total") or 0), cuerpo, None, der, cop)
        _llenar(prod, fila, 5, float(item.get("total") or 0) / total, cuerpo, None, der, "0.0%")
        fila += 1
    if fila == 4:
        _llenar(prod, 4, 1, "Sin productos en el período", mini, None, izq, merge=5)
        fila = 5
    _llenar(prod, fila, 1, "", etiqueta, fill_gray, cen)
    _llenar(prod, fila, 2, "TOTAL", etiqueta, fill_gray, izq)
    _llenar(prod, fila, 3, meta["items"], etiqueta, fill_gray, der)
    _llenar(prod, fila, 4, meta["total"], etiqueta, fill_gray, der, cop)
    _llenar(prod, fila, 5, 1, etiqueta, fill_gray, der, "0.0%")
    prod.freeze_panes = "A4"
    _pagina(prod)

    legal = wb.create_sheet("Marco legal")
    legal.column_dimensions["A"].width = 22
    legal.column_dimensions["B"].width = 88
    _llenar(legal, 1, 1, "Marco legal y notas para el contador", navy, fill_gold, izq, merge=2)
    for i, parrafo in enumerate(meta["marco_legal"], 3):
        legal.row_dimensions[i].height = 48
        _llenar(legal, i, 1, f"Nota {i - 2}", etiqueta, fill_gray, cen)
        _llenar(legal, i, 2, parrafo, cuerpo, None, izq)
    n = 3 + len(meta["marco_legal"])
    _llenar(legal, n, 1, "Normas de referencia", etiqueta, fill_gray, izq)
    _llenar(
        legal,
        n,
        2,
        "Código de Comercio arts. 48 y ss.; Estatuto Tributario (ingresos); facturación electrónica DIAN (resoluciones vigentes). Este POS no emite CUFE ni documento equivalente.",
        mini,
        None,
        izq,
    )
    legal.row_dimensions[n].height = 40
    _pagina(legal)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
